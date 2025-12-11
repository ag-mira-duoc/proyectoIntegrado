import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.utils import timezone
from .models import LogAuditoria
from usuarios.decorators import rol_requerido
import json

@login_required
@rol_requerido('Administradores', 'Auditores')
def exportar_logs_excel(request):
    """
    Genera un reporte Excel de los logs de auditoría aplicando los filtros actuales.
    """
    # 1. Obtener filtros de la petición GET
    usuario_id = request.GET.get('usuario')
    accion = request.GET.get('accion')
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')
    entidad = request.GET.get('entidad')

    # 2. Construir QuerySet base
    logs = LogAuditoria.objects.select_related('user').all().order_by('-fecha')

    # 3. Aplicar Filtros
    if usuario_id:
        logs = logs.filter(user__id=usuario_id)
    if accion:
        logs = logs.filter(accion=accion)
    if entidad:
        logs = logs.filter(tabla_afectada__icontains=entidad)
    
    # Filtro de fechas
    if fecha_inicio:
        logs = logs.filter(fecha__date__gte=fecha_inicio)
    if fecha_fin:
        logs = logs.filter(fecha__date__lte=fecha_fin)

    # 4. Crear el libro de Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Logs de Auditoría"

    # Estilos
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    alignment_center = Alignment(horizontal="center", vertical="center")
    alignment_wrap = Alignment(wrap_text=True, vertical="top")

    # Encabezados
    headers = [
        "Fecha/Hora", "Usuario", "Rol", "Acción", 
        "Entidad Afectada", "ID Registro", "Detalle", 
        "Valores Anteriores", "Valores Nuevos", "IP Origen"
    ]

    for col_num, column_title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = column_title
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = alignment_center

    # 5. Llenar filas
    for row_num, log in enumerate(logs, 2):
        # Formatear JSON para que sea legible
        val_ant = json.dumps(log.valores_anteriores, indent=2, ensure_ascii=False) if log.valores_anteriores else '-'
        val_nue = json.dumps(log.valores_nuevos, indent=2, ensure_ascii=False) if log.valores_nuevos else '-'
        
        # Obtener rol
        rol = "Usuario"
        if log.user:
            if log.user.es_administrador(): rol = "Administrador"
            elif log.user.es_analista(): rol = "Analista"
            elif log.user.es_auditor(): rol = "Auditor"

        fecha_local = timezone.localtime(log.fecha)
        fecha_str = fecha_local.strftime("%d/%m/%Y %H:%M:%S")

        row = [
            fecha_str,
            log.user.get_full_name() if log.user else "Sistema/Desconocido",
            rol,
            log.get_accion_display(),
            log.tabla_afectada or "-",
            str(log.registro_id) or "-",
            log.detalle,
            val_ant,
            val_nue,
            log.ip_address or "-"
        ]

        for col_num, cell_value in enumerate(row, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = cell_value
            cell.alignment = alignment_wrap

    # 6. Ajustar ancho de columnas
    column_widths = [20, 25, 15, 15, 20, 10, 40, 30, 30, 15]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    # 7. Generar respuesta HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    # Usar hora local también para el nombre del archivo
    timestamp = timezone.localtime(timezone.now()).strftime("%Y%m%d_%H%M")
    response['Content-Disposition'] = f'attachment; filename=Auditoria_NUAM_{timestamp}.xlsx'
    
    wb.save(response)
    return response